"""The same bounded operations build accommodation and equipment examples."""

import pytest
import yaml
from openpyxl import Workbook

from rangekeeper.graph.adapter import json as graph_json
from rangekeeper.graph.revision import Diff
from rangekeeper.graph.workflow import load, run, schema
from rangekeeper.graph.workflow.review import export


def example(tmp_path, domain="accommodation"):
    root = tmp_path / domain
    root.mkdir()
    inputs = root / "inputs"
    inputs.mkdir()
    spec = root / "spec"
    spec.mkdir()
    book = Workbook()
    sheet = book.active
    sheet.title = "Schedule"
    for row in [
        ("ID", "Level", "Size", "Label", "Other"),
        ("A1", "L1", 0, "2B", "2B"),
        ("A2", "L1", None, "2B", "3B"),
        ("A3", "L2", 12, "unknown", None),
        (None, None, None, None, None),
        ("note", None, None, None, None),
        ("END", None, None, None, None),
    ]:
        sheet.append(row)
    sheet["G1"] = "L1"
    sheet["G2"] = "L2"
    sheet["H1"] = 12
    book.save(inputs / "schedule.xlsx")

    def extract(id, start, end, columns):
        return {
            "id": id,
            "version": 1,
            "sheet": "Schedule",
            "rows": {"start": start, "end": end},
            "columns": [{"name": k, "column": v} for k, v in columns.items()],
        }

    source = {
        "version": 1,
        "namespace": "urn:example:" + domain,
        "steps": [
            {
                "id": "book",
                "operation": "read",
                "files": ["schedule.xlsx"],
                "source_key": "Schedule",
            },
            {
                "id": "raw",
                "operation": "extract",
                "input": "book",
                "unique_stop": True,
                "specification": {
                    "id": "raw",
                    "version": 1,
                    "sheet": "Schedule",
                    "rows": {
                        "start": 2,
                        "stop_before": {"column": "A", "equals": "END"},
                    },
                    "columns": [
                        {"name": k, "column": v}
                        for k, v in {
                            "code": "A",
                            "level": "B",
                            "size": "C",
                            "label": "D",
                            "other": "E",
                        }.items()
                    ],
                    "expect": {"cells": {"A1": "ID"}},
                },
            },
            {
                "id": "groups",
                "operation": "classify_rows",
                "input": "raw",
                "workbook": "book",
                "specification": {"identifier": "code", "pattern": "A[0-9]+"},
            },
            {
                "id": "selected",
                "operation": "select",
                "input": "groups",
                "where": {"column": "row_group", "equals": "matched"},
            },
            {
                "id": "numeric",
                "operation": "numbers",
                "input": "selected",
                "specifications": {
                    "number_size": {"column": "size", "nonnegative": True}
                },
            },
            {
                "id": "items",
                "operation": "transform",
                "input": "numeric",
                "specifications": {
                    "normalized": {
                        "operation": "normalize",
                        "columns": ["label"],
                        "case": "upper",
                    },
                    "label_count": {
                        "operation": "capture_integer",
                        "columns": ["normalized"],
                        "pattern": r"(\d+)B",
                        "group": 1,
                    },
                    "other_count": {
                        "operation": "capture_integer",
                        "columns": ["other"],
                        "pattern": r"(\d+)B",
                        "group": 1,
                    },
                    "count": {
                        "operation": "agreement",
                        "columns": ["label_count", "other_count"],
                    },
                    "product": {
                        "operation": "lookup",
                        "columns": ["normalized"],
                        "values": {"2B": "standard"},
                    },
                    "fallback_label": {
                        "operation": "fallback",
                        "columns": ["other", "label"],
                    },
                    "display": {
                        "operation": "format",
                        "columns": ["code"],
                        "template": "Item {code}",
                    },
                    "matched_label": {
                        "operation": "match",
                        "columns": ["normalized"],
                        "pattern": r"\d+B",
                    },
                },
            },
            {
                "id": "levels",
                "operation": "extract",
                "input": "book",
                "specification": extract("levels", 1, 2, {"code": "G"}),
            },
            {
                "id": "total_raw",
                "operation": "extract",
                "input": "book",
                "specification": extract("total", 1, 1, {"total": "H"}),
            },
            {
                "id": "total",
                "operation": "numbers",
                "input": "total_raw",
                "specifications": {"number": {"column": "total"}},
            },
        ],
    }
    model = {
        "version": 1,
        "taxonomy": {"code": "example", "name": domain},
        "classifications": [
            {"code": k, "name": k, "parent": None if k == "item" else "item"}
            for k in ("item", "level", "contains", "standard")
        ],
        "measures": [
            {
                "code": k,
                "name": k,
                "units": u,
                "quantity_kind": q,
                "aggregation": "NONE",
            }
            for k, u, q in [
                ("size", "meter**2", "AREA"),
                ("count", "dimensionless", "COUNT"),
            ]
        ],
        "templates": [
            {
                "id": "items",
                "table": "items",
                "kind": "entity",
                "identity_kind": "item",
                "key": {"column": "code"},
                "name": "Item {key}",
                "classification": "item",
                "decisions": ["D1"],
                "features": [{"name": "source_label", "binding": {"column": "label"}}],
                "measurements": [
                    {
                        "measure": "size",
                        "binding": {"column": "number_size"},
                        "on_unavailable": {"topic": "Unavailable measurement"},
                    },
                    {"measure": "count", "binding": {"column": "count"}},
                ],
                "labels": [{"name": "product", "bindings": [{"column": "product"}]}],
            },
            {
                "id": "levels",
                "table": "levels",
                "kind": "assembly",
                "identity_kind": "level",
                "key": {"column": "code"},
                "name": "Group {key}",
                "classification": "level",
            },
        ],
        "objects": [],
        "relationships": [
            {
                "id": "containment",
                "table": "items",
                "source": {"kind": "level", "key": {"column": "level"}},
                "target": {"kind": "item", "key": {"column": "code"}},
                "identity_kind": "contains",
                "key": "{source_key}:{target_key}",
                "classification": "contains",
            }
        ],
        "memberships": ["contains"],
    }
    checks = {
        "version": 1,
        "comparisons": [
            {
                "id": "population",
                "group": "Coverage",
                "scope": "Items",
                "left": {"kind": "graph_keys", "classification": "item"},
                "right": {"kind": "table_keys", "table": "items", "column": "code"},
                "purpose": "fidelity",
            },
            {
                "id": "total",
                "group": "Totals",
                "scope": "Sizes",
                "left": {
                    "kind": "graph_total",
                    "classification": "item",
                    "measure": "size",
                },
                "right": {
                    "kind": "evidence",
                    "binding": {"evidence": "total", "column": "number"},
                },
                "purpose": "reconciliation",
            },
            {
                "id": "fidelity",
                "group": "Values",
                "scope": "{key}",
                "each": "items",
                "scope_column": "code",
                "left": {
                    "kind": "graph_measurement",
                    "identity_kind": "item",
                    "key": {"column": "code"},
                    "measure": "size",
                },
                "right": {"kind": "column", "binding": {"column": "number_size"}},
                "purpose": "fidelity",
            },
        ],
        "invariants": ["fact_coverage", "membership", "defined_kinds"],
        "deferred": ["No allocation is inferred"],
    }
    decisions = {
        "version": 1,
        "decisions": [
            {
                "id": "D1",
                "status": "accepted",
                "text": "Use explicit observations; unknown suffix remains unknown.",
                "source": "review",
                "date": "2026-01-01",
            }
        ],
    }
    documents = {
        "sources": source,
        "model": model,
        "checks": checks,
        "decisions": decisions,
    }
    for k, v in documents.items():
        (spec / (k + ".yaml")).write_text(yaml.safe_dump(v, sort_keys=False))
    return root, documents


@pytest.mark.parametrize("domain", ["accommodation", "equipment"])
def test_synthetic_vertical_slice(tmp_path, domain):
    root, _ = example(tmp_path, domain)
    spec = load(root / "spec")
    result = run(spec, input_root=root / "inputs")
    assert result.output is not None, result.diagnostics
    built = result.output
    assert len(built.graph.entities) == 5 and len(built.graph.relationships) == 3
    objects = {e.code: e for e in built.graph.entities}
    assert objects["A1"].measurements["size"].quantity.magnitude == 0
    assert "count" not in objects["A2"].measurements
    assert "count" not in objects["A3"].measurements
    assert any(i.code == "conflicting_values" for i in built.evidence["items"].issues)
    assert [r.values["row_group"] for r in built.evidence["groups"].data.rows] == [
        "matched",
        "matched",
        "matched",
        "blank",
        "other",
    ]
    assert next(c for c in built.checks if c.id == "total").status == "unavailable"
    assert next(c for c in built.checks if c.id == "total").known_subtotal == 12
    encoded = graph_json.dumps(built.graph)
    assert not Diff.between(built.graph, graph_json.loads(encoded)).changed
    assert encoded == graph_json.dumps(
        run(spec, input_root=root / "inputs").output.graph
    )
    assert built.operations and all(o.method.version for o in built.operations)
    assert not (root / "artifacts").exists()
    export(built, root / "artifacts")
    assert (root / "artifacts/viewer.html").is_file()
    assert "unavailable" in (root / "artifacts/review.html").read_text()


@pytest.mark.parametrize(
    "mutation",
    [
        lambda d: d["sources"]["steps"][1].update(input="later"),
        lambda d: d["sources"]["steps"][0].update(operation="python"),
        lambda d: d["sources"]["steps"][0].update(callable="os.system"),
        lambda d: d["sources"]["steps"].append(d["sources"]["steps"][0]),
        lambda d: d["model"]["templates"][0].update(arbitrary="expression"),
        lambda d: d["sources"]["steps"][5]["specifications"]["display"].update(
            template="{code.__class__}"
        ),
    ],
)
def test_malformed_specs_fail_before_run(tmp_path, mutation):
    root, docs = example(tmp_path)
    mutation(docs)
    for k, v in docs.items():
        (root / "spec" / f"{k}.yaml").write_text(yaml.safe_dump(v, sort_keys=False))
    with pytest.raises((ValueError, TypeError)):
        load(root / "spec")


@pytest.mark.parametrize(
    "content",
    [
        "version: 1\nversion: 1\n",
        "!!python/object/apply:os.system [echo bad]",
        "version: 1\nx: &a [*a]",
    ],
)
def test_unsafe_yaml_rejected(tmp_path, content):
    root, _ = example(tmp_path)
    (root / "spec/sources.yaml").write_text(content)
    with pytest.raises((ValueError, yaml.YAMLError)):
        load(root / "spec")


def test_unavailable_input_and_duplicate_keys(tmp_path):
    root, _docs = example(tmp_path)
    (root / "inputs/schedule.xlsx").unlink()
    assert run(load(root / "spec"), input_root=root / "inputs").output is None
    assert schema()["executableContent"] is False


def rewrite(root, docs):
    for key, value in docs.items():
        (root / "spec" / f"{key}.yaml").write_text(
            yaml.safe_dump(value, sort_keys=False)
        )


def membership_check():
    return {
        "id": "members",
        "group": "Membership",
        "scope": "{key}",
        "each": "levels",
        "scope_column": "code",
        "report": "counts",
        "left": {
            "kind": "membership_keys",
            "identity_kind": "level",
            "key": {"column": "code"},
            "classification": "item",
        },
        "right": {
            "kind": "table_keys",
            "table": "items",
            "column": "code",
            "where": {"column": "level", "binding": {"column": "code"}},
        },
    }


def test_equal_counts_do_not_hide_wrong_members(tmp_path):
    root, docs = example(tmp_path)
    docs["sources"]["steps"][5]["specifications"]["rewired"] = {
        "operation": "lookup",
        "columns": ["code"],
        "values": {"A1": "L2", "A2": "L1", "A3": "L1"},
    }
    docs["model"]["relationships"][0]["source"]["key"] = {"column": "rewired"}
    docs["checks"]["comparisons"].append(membership_check())
    rewrite(root, docs)
    result = run(load(root / "spec"), input_root=root / "inputs").output
    members = [c for c in result.checks if c.group == "Membership"]
    assert all(c.left == c.right and c.status == "difference" for c in members)
    assert all(c.left_members != c.right_members for c in members)


def test_membership_checks_use_final_assemblies(tmp_path):
    root, docs = example(tmp_path)
    docs["checks"]["comparisons"].append(membership_check())
    rewrite(root, docs)
    result = run(load(root / "spec"), input_root=root / "inputs").output
    members = [c for c in result.checks if c.group == "Membership"]
    assert [(c.left, c.right, c.status) for c in members] == [
        (2, 2, "agree"),
        (1, 1, "agree"),
    ]


@pytest.mark.parametrize(
    "case,code",
    [
        ("duplicate", "duplicate_business_key"),
        ("unmatched", "unmatched_relationship"),
        ("collision", "output_collision"),
        ("missing_column", "missing_column"),
        ("cycle", "membership_cycle"),
    ],
)
def test_valid_requests_with_incompatible_sources_return_unavailable(
    tmp_path, case, code
):
    root, docs = example(tmp_path)
    if case == "duplicate":
        from openpyxl import load_workbook

        path = root / "inputs/schedule.xlsx"
        book = load_workbook(path)
        book.active["A3"] = "A1"
        book.save(path)
    elif case == "unmatched":
        docs["model"]["relationships"][0]["source"]["key"] = {"value": "unknown"}
    elif case == "collision":
        docs["sources"]["steps"][4]["specifications"]["size"] = {"column": "size"}
    elif case == "missing_column":
        docs["sources"]["steps"][4]["specifications"]["number_size"] = {
            "column": "absent"
        }
    else:
        docs["model"]["relationships"].extend([
            {
                "id": f"cycle-{left}",
                "source": {"kind": "level", "key": {"value": left}},
                "target": {"kind": "level", "key": {"value": right}},
                "identity_kind": "contains",
                "key": "{source_key}:{target_key}",
                "classification": "contains",
            }
            for left, right in [("L1", "L2"), ("L2", "L1")]
        ])
    rewrite(root, docs)
    outcome = run(load(root / "spec"), input_root=root / "inputs")
    assert outcome.output is None
    assert outcome.diagnostics[0].code == code


def test_formula_without_cache_is_occupied_and_marker_must_be_unique(tmp_path):
    from openpyxl import load_workbook

    root, _docs = example(tmp_path)
    path = root / "inputs/schedule.xlsx"
    book = load_workbook(path)
    book.active["C5"] = "=1+2"
    book.save(path)
    outcome = run(load(root / "spec"), input_root=root / "inputs")
    assert outcome.output is not None
    groups = outcome.output.evidence["groups"]
    assert groups.data.rows[3].values["row_group"] == "other"
    assert any(i.code == "missing_formula_cache" for i in groups.issues)
    book.active["A8"] = "END"
    book.save(path)
    outcome = run(load(root / "spec"), input_root=root / "inputs")
    assert (
        outcome.output is None
        and outcome.diagnostics[0].code == "nonunique_stopping_marker"
    )


def test_shared_membership_does_not_double_count_atomic_total(tmp_path):
    root, docs = example(tmp_path)
    docs["model"]["relationships"].append({
        "id": "shared",
        "source": {"kind": "level", "key": {"value": "L2"}},
        "target": {"kind": "item", "key": {"value": "A1"}},
        "identity_kind": "contains",
        "key": "{source_key}:{target_key}",
        "classification": "contains",
    })
    rewrite(root, docs)
    result = run(load(root / "spec"), input_root=root / "inputs").output
    objects = {x.code: x for x in result.graph.entities}
    assert objects["A1"].id in objects["L1"].entity_ids & objects["L2"].entity_ids
    assert next(c for c in result.checks if c.id == "total").known_subtotal == 12


def test_business_identity_independent_of_row_edition_and_value(tmp_path):
    from openpyxl import load_workbook

    root, _docs = example(tmp_path)
    spec = load(root / "spec")
    before = run(spec, input_root=root / "inputs").output.graph
    path = root / "inputs/schedule.xlsx"
    book = load_workbook(path)
    sheet = book.active
    row2 = [sheet.cell(2, c).value for c in range(1, 6)]
    row4 = [sheet.cell(4, c).value for c in range(1, 6)]
    for c, value in enumerate(row4, 1):
        sheet.cell(2, c, value)
    for c, value in enumerate(row2, 1):
        sheet.cell(4, c, value)
    sheet["C4"] = 1
    book.save(path)
    after = run(spec, input_root=root / "inputs").output.graph
    assert {x.code: x.id for x in before.entities} == {
        x.code: x.id for x in after.entities
    }
    assert {x.id for x in before.relationships} == {x.id for x in after.relationships}
