"""Capture XLSX bytes, decode native observations, then construct an immutable snapshot."""

import posixpath
import re
from collections.abc import Iterator
from contextlib import ExitStack, contextmanager
from hashlib import sha256
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING
from uuid import UUID, uuid5
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from .... import validate
from ...provenance import Location, Method, Source
from .. import _structured
from ..operation import Operation, Outcome, _Failure, _invoke, fingerprint
from ._coordinates import address
from .snapshot import Cell, Workbook, Worksheet

if TYPE_CHECKING:
    from openpyxl.workbook.workbook import Workbook as NativeWorkbook
    from openpyxl.worksheet.worksheet import Worksheet as NativeWorksheet

_NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_METHOD = Method(code="rk.excel.read", version="1")


def _xml_sheets(data: bytes, source: Source) -> dict[str, dict[str, ET.Element]]:
    """Retain XML observations that formatted openpyxl values cannot represent."""
    with ZipFile(BytesIO(data)) as archive:
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item for item in relationships}
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        result = {}
        for sheet in book.findall("s:sheets/s:sheet", _NS):
            relationship = targets[sheet.attrib[_REL]]
            relationship_type = relationship.attrib.get("Type", "")
            if not relationship_type.endswith("/worksheet"):
                raise _Failure(
                    "unsupported_workbook",
                    "Only worksheet content is supported in this reader",
                    locations=(Location(source=source),),
                    details={
                        "sheet": sheet.attrib["name"],
                        "relationship_type": relationship_type,
                    },
                )
            target = relationship.attrib["Target"]
            member = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            xml = ET.fromstring(archive.read(member))
            cells = {}
            for cell in xml.findall("s:sheetData/s:row/s:c", _NS):
                coordinate = cell.attrib["r"]
                address(coordinate)
                if coordinate in cells:
                    raise ValueError(f"Duplicate physical cell {coordinate}")
                cells[coordinate] = cell
            result[sheet.attrib["name"]] = cells
        return result


@contextmanager
def _opened(
    data: bytes,
    source: Source,
) -> Iterator[
    tuple["NativeWorkbook", "NativeWorkbook", dict[str, dict[str, ET.Element]]]
]:
    """Translate decoding errors only; adapter construction defects must propagate."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    with ExitStack() as resources:
        try:
            xml = _xml_sheets(data, source)
            formulas = load_workbook(BytesIO(data), data_only=False, keep_links=False)
            resources.callback(formulas.close)
            values = load_workbook(BytesIO(data), data_only=True, keep_links=False)
            resources.callback(values.close)
        except (
            BadZipFile,
            ET.ParseError,
            InvalidFileException,
            KeyError,
            ValueError,
            IndexError,
            OSError,
        ) as exc:
            raise _Failure(
                "invalid_workbook",
                "Workbook could not be decoded",
                locations=(Location(source=source),),
                details={"reason": str(exc), "error_type": type(exc).__name__},
            ) from exc
        # Keep yield outside the decoding try: caller defects are not source failures.
        yield formulas, values, xml


def _read_cell(
    sheet: "NativeWorksheet",
    values: "NativeWorksheet",
    coordinate: str,
    element: ET.Element,
    source: Source,
) -> Cell:
    native = sheet[coordinate]
    cached_cell = values[coordinate]
    location = Location(
        source=source,
        reference={"sheet": sheet.title, "cell": coordinate},
    )
    formula_element = element.find("s:f", _NS)
    value_element = element.find("s:v", _NS)
    xml_type = element.get("t")
    formula = None
    if formula_element is not None:
        formula = (
            native.value
            if type(native.value) is str
            else getattr(native.value, "text", None)
        )
        if type(formula) is not str:
            raise _Failure(
                "unsupported_formula",
                "Formula kind cannot be represented faithfully",
                locations=(location,),
                details={
                    "formula_attributes": tuple(sorted(formula_element.attrib.items()))
                },
            )

    # Numeric empty <v/> is not a cache. String <v/> explicitly stores "".
    cache_present = (
        formula is not None
        and value_element is not None
        and (value_element.text is not None or xml_type == "str")
    )
    cached = cached_cell.value if formula is not None else None
    if cache_present and xml_type == "str" and cached is None:
        cached = ""
    raw = formula if formula is not None else native.value
    if formula is None and raw is None and xml_type in ("inlineStr", "str"):
        raw = ""
    return Cell(
        location=location,
        raw=raw,
        formula=formula,
        cached=cached,
        cache_present=cache_present,
        data_type=native.data_type,
        cached_type=cached_cell.data_type if formula is not None else None,
        number_format=native.number_format,
        xml_value=value_element.text if value_element is not None else None,
        xml_type=xml_type,
        xml_formula=formula_element.text if formula_element is not None else None,
        formula_attributes=(
            tuple(sorted(formula_element.attrib.items()))
            if formula_element is not None
            else ()
        ),
    )


def _read_sheet(
    sheet: "NativeWorksheet",
    values: "NativeWorksheet",
    xml: dict[str, ET.Element],
    source: Source,
) -> Worksheet:
    cells = {
        coordinate: _read_cell(sheet, values, coordinate, element, source)
        for coordinate, element in xml.items()
    }
    return Worksheet(
        location=Location(source=source, reference={"sheet": sheet.title}),
        state=sheet.sheet_state,
        declared_rows=sheet.max_row,
        declared_columns=sheet.max_column,
        cells=cells,
        merged_ranges=tuple(sorted(str(item) for item in sheet.merged_cells)),
    )


def _parse(data: bytes, source: Source, operation: Operation) -> Workbook:
    with _opened(data, source) as (formulas, values, xml):
        sheets = tuple(
            _read_sheet(sheet, values[sheet.title], xml[sheet.title], source)
            for sheet in formulas.worksheets
        )
        return Workbook(
            source=source,
            worksheets=sheets,
            reader_fingerprint=fingerprint(operation),
            metadata={
                "created": formulas.properties.created,
                "modified": formulas.properties.modified,
                "epoch": formulas.epoch,
                "parser": operation.specification["parser"],
            },
        )


def read(
    path: str | Path,
    *,
    namespace: UUID,
    source_key: str,
    name: str,
    expected_checksum: str | None = None,
) -> Outcome[Workbook]:
    """Read a local .xlsx edition. Paths never contribute to semantic identity."""
    if not isinstance(path, (str, Path)):
        raise TypeError("path must be str or Path")
    validate.require_uuid(namespace, "namespace")
    validate.require_text(source_key, "source_key")
    validate.require_text(name, "name")
    if expected_checksum is not None:
        if (
            type(expected_checksum) is not str
            or re.fullmatch(r"(?:sha256:)?[0-9a-fA-F]{64}", expected_checksum) is None
        ):
            raise ValueError("expected_checksum must be a SHA-256 hex digest")
        expected_checksum = expected_checksum.removeprefix("sha256:").lower()
    path = Path(path)
    try:
        parser_version = version("openpyxl")
    except PackageNotFoundError:
        parser_version = None
    specification = {
        "namespace": namespace,
        "source_key": source_key,
        "name": name,
        "expected_checksum": expected_checksum,
        "format": "xlsx",
        "parser": {"name": "openpyxl", "version": parser_version},
        "formula_values": "preserve",
        "external_links": "do_not_load",
    }
    data = None
    problem = None
    if path.suffix.lower() != ".xlsx":
        problem = _Failure(
            "unsupported_format",
            "Expected an .xlsx file",
            details={"suffix": path.suffix},
        )
    else:
        try:
            data = path.read_bytes()
        except OSError as exc:
            problem = _Failure(
                "source_unavailable",
                "Source file could not be read",
                details={
                    "source_key": source_key,
                    "path": str(path),
                    "errno": exc.errno,
                },
            )
    checksum = sha256(data).hexdigest() if data is not None else None

    def execute(operation: Operation) -> Workbook:
        if problem is not None:
            raise problem
        assert data is not None and checksum is not None
        source = Source(
            id=uuid5(
                namespace,
                _structured.fingerprint((
                    "rk.excel.source/v1",
                    source_key,
                    checksum,
                )),
            ),
            name=name,
            checksum=checksum,
        )
        if expected_checksum is not None and checksum != expected_checksum:
            raise _Failure(
                "checksum_mismatch",
                "Source edition differs from the expected checksum",
                locations=(Location(source=source),),
                details={"expected": expected_checksum, "actual": checksum},
            )
        if parser_version is None:
            raise _Failure(
                "dependency_unavailable",
                "Install Rangekeeper[excel] to read XLSX",
                locations=(Location(source=source),),
                details={"dependency": "openpyxl"},
            )
        return _parse(data, source, operation)

    return _invoke(
        _METHOD,
        specification,
        {"source": "sha256:" + checksum if checksum is not None else None},
        execute,
    )
